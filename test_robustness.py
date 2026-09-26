"""空间查询边界与 Excel 文本输入的回归测试。"""

from __future__ import annotations

import io
import math
import random
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from cable_stat.geometry import SegmentGrid, project_to_segment
from cable_stat.loaders import load_workbook_rows
from cable_stat.models import Room, Segment
from cable_stat.rooms import RoomAnalyzer, rooms_for_point
from cable_stat.text import normalize_cabinet_name, normalize_text


def make_segment(idx: int, ax: float, ay: float, bx: float, by: float) -> Segment:
    length = math.hypot(bx - ax, by - ay)
    return Segment(idx, f"R{idx}", "1F", ax, ay, bx, by, length, length, "", "", "")


class BoundedCellLookup(dict):
    """用查找次数验证工作量边界，避免依赖机器速度的耗时断言。"""

    def __init__(self, cells):
        super().__init__(cells)
        self.lookups = 0

    def get(self, key, default=None):
        self.lookups += 1
        if self.lookups > 20_000:
            raise AssertionError("最近线段查询扫描了过多空格，未及时退化为全量比较")
        return super().get(key, default)


class SegmentGridRobustnessTests(unittest.TestCase):
    def test_far_away_query_has_bounded_work_and_correct_tie_break(self):
        segments = [make_segment(9, 0, 1, 10, 1), make_segment(2, 0, -1, 10, -1)]
        grid = SegmentGrid(segments)
        grid.cells = BoundedCellLookup(grid.cells)
        result = grid.nearest(1e8, 0)
        self.assertIsNotNone(result)
        self.assertEqual(result[1], 2)
        self.assertEqual(result[0], math.hypot(1e8 - 10, 1))
        self.assertEqual(result[3:], (10, -1, 10))

    def test_sparse_drawing_query_inside_huge_bounds(self):
        segments = [make_segment(0, -1e8, 0, -1e8 + 1, 0), make_segment(1, 1e8, 0, 1e8 + 1, 0)]
        grid = SegmentGrid(segments)
        grid.cells = BoundedCellLookup(grid.cells)
        result = grid.nearest(0, 0)
        self.assertEqual(result[1], 0)
        self.assertEqual(result[0], 1e8 - 1)

    def test_index_matches_full_scan_for_local_and_distant_points(self):
        rng = random.Random(270926)
        segments = [
            make_segment(i, rng.uniform(-50, 50), rng.uniform(-50, 50), rng.uniform(-50, 50), rng.uniform(-50, 50))
            for i in range(35)
        ]
        grid = SegmentGrid(segments, pad=0.05)
        points = [(rng.uniform(-75, 75), rng.uniform(-75, 75)) for _ in range(100)]
        points.extend([(1e8, -1e8), (-1e8, 1e8), (1e8, 1e8)])
        for x, y in points:
            expected = min((project_to_segment(x, y, seg)[3], seg.idx) for seg in segments)
            self.assertEqual(grid.nearest(x, y)[:2], expected)

    def test_empty_index(self):
        self.assertIsNone(SegmentGrid([]).nearest(1e8, 1e8))


class RoomBoundaryRobustnessTests(unittest.TestCase):
    def setUp(self):
        self.room = Room("保护室", "1F", [(0, 0), (10, 0), (10, 10), (0, 10)], "", "")

    def test_bbox_preserves_polygon_boundary_tolerance_on_every_side(self):
        points = [(-5e-9, 5), (10 + 5e-9, 5), (5, -5e-9), (5, 10 + 5e-9)]
        graph = SimpleNamespace(nodes={str(i): {"floor": "1F", "x": x, "y": y} for i, (x, y) in enumerate(points)})
        analyzer = RoomAnalyzer(graph, [self.room])
        for i, (x, y) in enumerate(points):
            expected = rooms_for_point("1F", x, y, [self.room])
            self.assertEqual(expected, [self.room])
            self.assertEqual(analyzer.rooms_at(str(i)), expected)

    def test_bbox_margin_does_not_bypass_polygon_or_floor_test(self):
        graph = SimpleNamespace(nodes={
            "outside": {"floor": "1F", "x": -5e-7, "y": 5},
            "other_floor": {"floor": "2F", "x": 5, "y": 5},
        })
        analyzer = RoomAnalyzer(graph, [self.room])
        self.assertEqual(analyzer.rooms_at("outside"), [])
        self.assertEqual(analyzer.rooms_at("other_floor"), [])


class TextRobustnessTests(unittest.TestCase):
    def test_illegal_controls_are_removed_before_excel_roundtrip(self):
        controls = "".join(chr(i) for i in range(32) if i not in (9, 10, 13))
        normalized = normalize_text("保护" + controls + "屏①")
        self.assertEqual(normalized, "保护屏①")
        workbook = Workbook()
        workbook.active["A1"] = normalized
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        buffer.seek(0)
        restored = load_workbook(buffer)
        self.assertEqual(restored.active["A1"].value, "保护屏①")
        restored.close()

    def test_legal_whitespace_and_normal_text_keep_existing_behavior(self):
        self.assertEqual(normalize_text("  保护\t屏\nA\r\n　１  "), "保护 屏 A １")
        self.assertEqual(normalize_cabinet_name(" Ａ\x07柜\n １ "), "A柜1")
        self.assertEqual(normalize_text(None), "")
        self.assertEqual(normalize_text(42), "42")


class WorkbookRowRobustnessTests(unittest.TestCase):
    def load_rows(self, workbook):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "清册.xlsx"
            workbook.save(path)
            loaded, _, _, rows, names = load_workbook_rows(path)
            loaded.close()
        workbook.close()
        return rows, names

    def test_merged_multiline_header_and_repeated_headers_are_not_cables(self):
        workbook = Workbook()
        ws = workbook.active
        ws.append(["电缆编号", "起点", "终点", "电缆长度"])
        ws.append([None, None, None, "m"])
        for column in "ABC":
            ws.merge_cells(f"{column}1:{column}2")
        ws.append(["C1", "A柜", "B柜", 12])
        ws.append(["电缆编号", "起\n点", "终 点", "电缆长度"])
        ws.append([None, "起点", "终点", None])
        ws.append(["C2", "C柜", "D柜", 8])
        rows, names = self.load_rows(workbook)
        self.assertEqual([row["row_no"] for row in rows], [3, 6])
        self.assertEqual(names, ["A柜", "B柜", "C柜", "D柜"])

    def test_horizontal_notes_are_skipped_but_incomplete_cables_remain(self):
        workbook = Workbook()
        ws = workbook.active
        ws.append(["电缆编号", "起点", "终点", "电缆长度", "备注"])
        ws.append(["一、控制电缆"])
        ws.merge_cells("A2:E2")
        ws.append([None, "说明：长度以现场为准"])
        ws.merge_cells("B3:E3")
        ws.append(["C1", None, "B柜", 12])
        ws.append(["C2", "A柜", None, 8])
        ws.append(["C3", None, None, 8])
        ws.append(["C4", "A柜", "B柜", 12, "检查备注"])
        rows, names = self.load_rows(workbook)
        self.assertEqual([row["电缆编号"] for row in rows], ["C1", "C2", "C3", "C4"])
        self.assertEqual([(row["起点"], row["终点"]) for row in rows[:3]], [("", "B柜"), ("A柜", ""), ("", "")])
        self.assertEqual(names, ["A柜", "B柜"])

    def test_vertical_merged_endpoints_and_horizontal_nonkey_columns_remain(self):
        workbook = Workbook()
        ws = workbook.active
        ws.append(["电缆编号", "起点", "终点", "电缆长度", "备注", "校核"])
        ws.append(["C1", "A柜", "B柜", 10, "备注"])
        ws.append(["C2", None, None, 8])
        ws.merge_cells("B2:B3")
        ws.merge_cells("C2:C3")
        ws.merge_cells("E2:F2")
        rows, names = self.load_rows(workbook)
        self.assertEqual([row["电缆编号"] for row in rows], ["C1", "C2"])
        self.assertEqual([(row["起点"], row["终点"]) for row in rows], [("A柜", "B柜"), ("A柜", "B柜")])
        self.assertEqual(names, ["A柜", "B柜"])


if __name__ == "__main__":
    unittest.main()
