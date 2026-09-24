from __future__ import annotations

import sys
import traceback
from pathlib import Path

import openpyxl

import calculate_cable_lengths
from calculate_cable_lengths import normalize_header


# 与 load_workbook_rows 的必需表头一致；「自动统计」「向上取整」列可以没有，核心脚本会自动补建
REQUIRED_HEADERS = {"电缆编号", "起点", "终点", "电缆长度"}


def has_required_headers(path: Path) -> bool:
    if path.name.startswith("~$"):
        return False
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        ws = wb.active
        headers = {normalize_header(cell.value) for cell in ws[1] if cell.value is not None}
        wb.close()
        return REQUIRED_HEADERS.issubset(headers)
    except Exception:
        return False


def find_workbook(project_dir: Path) -> Path:
    candidates = [path for path in project_dir.glob("*.xlsx") if has_required_headers(path)]
    if not candidates:
        raise FileNotFoundError(
            f"没有找到表头包含“电缆编号、起点、终点、电缆长度”的 xlsx 文件（表头空格/换行会自动忽略，“自动统计”列可以没有）。\n"
            f"请把原清册（建议命名为“自动统计.xlsx”）放到项目文件夹：{project_dir}"
        )
    exact = [path for path in candidates if path.stem == "自动统计"]
    if exact:
        return exact[0]
    if len(candidates) == 1:
        return candidates[0]
    names = "\n".join(str(path) for path in candidates)
    raise RuntimeError(f"找到多个候选工作簿，请只保留一个自动统计表或改名为“自动统计.xlsx”：\n{names}")


def main() -> int:
    if getattr(sys, "frozen", False):
        tool_dir = Path(sys.executable).resolve().parent
    else:
        tool_dir = Path(__file__).resolve().parent
    if len(sys.argv) > 1:
        # 兼容 cmd 里 "%~dp0" 这类尾反斜杠把引号转义进参数的情况（引号不可能出现在合法路径里）
        raw = sys.argv[1].strip().strip('"').rstrip("\\")
        project_dir = Path(raw).resolve() if raw else Path.cwd().resolve()
    else:
        project_dir = Path.cwd().resolve()
    if project_dir == tool_dir:
        print("这里是通用工具文件夹，不是项目文件夹。")
        print("请双击某个项目文件夹里的 一键启动.cmd 来运行统计。")
        print("新项目：双击本文件夹的 新建项目.cmd，或复制“项目模板”文件夹到任何地方再改名。")
        return 1
    workbook = find_workbook(project_dir)
    output = project_dir / "outputs" / "自动统计_计算结果.xlsx"
    data_dir = project_dir / "data"

    try:
        calculate_cable_lengths.main(
            [
                "--workbook",
                str(workbook),
                "--data-dir",
                str(data_dir),
                "--output",
                str(output),
            ]
        )
    except Exception as exc:
        traceback.print_exc()
        print(f"\n计算失败：{exc}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
