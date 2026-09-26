"""calculate_cable_lengths.py 的轻量测试。

运行方式（任选其一）：
    python -m pytest test_calculate_cable_lengths.py -q
    python test_calculate_cable_lengths.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from calculate_cable_lengths import (
    Cabinet,
    Room,
    RouteGraph,
    Segment,
    Shaft,
    apply_aliases,
    build_graph,
    calculate_rows,
    compute_floor_elevations,
    compute_floor_offsets,
    is_route_bent,
    load_aliases,
    load_rule_overrides,
    load_rooms,
    load_params,
    normalize_cabinet_name,
    normalize_shaft_base_id,
    shaft_suffix_floor,
    point_in_polygon,
    segment_room_relation,
)


PARAMS = {
    "竖井高度": 4.5,
    "不拐弯修正": 7.0,
    "拐弯倍率": 1.2,
    "CAD每米单位": 1000.0,
    "吸附容差": 10.0,
    "最大接入距离": 20000.0,
    "跨房间修正": 3.0,
}


def make_segment(idx: int, ax: float, ay: float, bx: float, by: float, floor: str = "1F") -> Segment:
    import math

    length = math.hypot(bx - ax, by - ay)
    return Segment(
        idx=idx,
        route_id=f"R{idx}",
        floor=floor,
        ax=ax,
        ay=ay,
        bx=bx,
        by=by,
        length_units=length,
        length_m=length / PARAMS["CAD每米单位"],
        layer=f"CABLE_ROUTE_{floor}",
        handle=f"H{idx}",
        section_no="1",
    )


def make_cabinet(name: str, x: float, y: float, floor: str = "1F") -> Cabinet:
    return Cabinet(name=name, floor=floor, x=x, y=y, layer=f"CABLE_CABINET_{floor}", object_type="TEXT", handle="")


def make_shaft(shaft_id: str, x: float, y: float, floor: str) -> Shaft:
    return Shaft(
        shaft_id=shaft_id,
        base_id=normalize_shaft_base_id(shaft_id),
        floor=floor,
        x=x,
        y=y,
        height_m=PARAMS["竖井高度"],
        layer=f"CABLE_SHAFT_{floor}",
        object_type="TEXT",
        handle="",
    )


def make_rows(pairs: list[tuple[str, str]]) -> list[dict]:
    return [
        {"row_no": i + 2, "电缆编号": f"C{i + 1}", "起点": a, "终点": b, "手工长度": None}
        for i, (a, b) in enumerate(pairs)
    ]


def make_room(name: str, vertices: list[tuple[float, float]], floor: str = "1F") -> Room:
    return Room(name=name, floor=floor, vertices=vertices, layer=f"CABLE_ROOM_{floor}", handle=name)


def calculate_pair_with_rooms(
    segments: list[Segment], cabinets: dict[str, Cabinet], rooms: list[Room], shafts: list[Shaft] | None = None,
    overrides=None, params=None,
):
    params = params or PARAMS
    graph, nodes, _ = build_graph(segments, cabinets, shafts or [], list(cabinets), params)
    details, visuals = calculate_rows(make_rows([("柜A", "柜B")]), graph, nodes, cabinets, params, overrides, rooms)
    return details[0], visuals[0]


def test_same_room_does_not_add_room_extra():
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 0), "柜B": make_cabinet("柜B", 9000, 0)}
    room = make_room("配电室", [(-1000, -1000), (11000, -1000), (11000, 1000), (-1000, 1000)])
    row, visual = calculate_pair_with_rooms(segments, cabinets, [room])
    assert row["自动统计"] == 15.0, row
    assert row["跨房间"] == "否" and visual["crosses_room"] is False
    assert row["起点房间"] == "配电室" and row["终点房间"] == "配电室"


def test_different_rooms_add_room_extra_once():
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 0), "柜B": make_cabinet("柜B", 9000, 0)}
    rooms = [
        make_room("房间A", [(-500, -1000), (4000, -1000), (4000, 1000), (-500, 1000)]),
        make_room("房间B", [(6000, -1000), (10500, -1000), (10500, 1000), (6000, 1000)]),
    ]
    row, _ = calculate_pair_with_rooms(segments, cabinets, rooms)
    assert row["自动统计"] == 18.0, row
    assert row["修正量"] == 10.0 and "跨房间 +3" in row["规则"]


def test_same_room_endpoints_but_route_leaves_room():
    segments = [
        make_segment(1, 1000, 1000, 1000, 8000),
        make_segment(2, 1000, 8000, 9000, 8000),
        make_segment(3, 9000, 8000, 9000, 1000),
    ]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 1000), "柜B": make_cabinet("柜B", 9000, 1000)}
    room = make_room("控制室", [(0, 0), (10000, 0), (10000, 5000), (0, 5000)])
    row, _ = calculate_pair_with_rooms(segments, cabinets, [room])
    assert row["起点房间"] == "控制室" and row["终点房间"] == "控制室"
    assert row["跨房间"] == "是", row


def test_room_to_outside_and_outside_only():
    segments = [make_segment(1, 0, 0, 10000, 0)]
    room = make_room("房间A", [(-500, -1000), (4000, -1000), (4000, 1000), (-500, 1000)])
    cabinets = {"柜A": make_cabinet("柜A", 1000, 0), "柜B": make_cabinet("柜B", 9000, 0)}
    row, _ = calculate_pair_with_rooms(segments, cabinets, [room])
    assert row["跨房间"] == "是"
    outdoor_room = make_room("远处房间", [(20000, 20000), (21000, 20000), (21000, 21000), (20000, 21000)])
    row2, _ = calculate_pair_with_rooms(segments, cabinets, [outdoor_room])
    assert row2["跨房间"] == "否" and row2["自动统计"] == 15.0


def test_outside_path_through_room_adds_extra():
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {"柜A": make_cabinet("柜A", 500, 0), "柜B": make_cabinet("柜B", 9500, 0)}
    room = make_room("中间房间", [(4000, -1000), (6000, -1000), (6000, 1000), (4000, 1000)])
    row, _ = calculate_pair_with_rooms(segments, cabinets, [room])
    assert not row["起点房间"] and not row["终点房间"]
    assert row["跨房间"] == "是"


def test_cross_floor_rooms_add_extra():
    segments = [make_segment(1, 0, 0, 10000, 0, "1F"), make_segment(2, 0, 0, 10000, 0, "2F")]
    shafts = [make_shaft("ZJ1_1F", 5000, 0, "1F"), make_shaft("ZJ1_2F", 5000, 0, "2F")]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 0, "1F"), "柜B": make_cabinet("柜B", 9000, 0, "2F")}
    bounds = [(-1000, -1000), (11000, -1000), (11000, 1000), (-1000, 1000)]
    rooms = [make_room("一层房间", bounds, "1F"), make_room("二层房间", bounds, "2F")]
    row, _ = calculate_pair_with_rooms(segments, cabinets, rooms, shafts)
    assert row["自动统计"] == 26.4, row
    assert row["跨房间"] == "是"


def test_room_extra_stacks_after_override():
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 0), "柜B": make_cabinet("柜B", 9000, 0)}
    rooms = [make_room("房间A", [(-500, -1000), (4000, -1000), (4000, 1000), (-500, 1000)])]
    overrides = {("柜A", "柜B"): {"extra": 2.0, "note": "实测"}}
    row, _ = calculate_pair_with_rooms(segments, cabinets, rooms, overrides=overrides)
    assert row["自动统计"] == 13.0 and row["修正量"] == 5.0, row
    assert "强制规则 +2" in row["规则"] and "跨房间 +3" in row["规则"]


def test_boundary_and_concave_polygon_geometry():
    concave = make_room("凹房间", [(0, 0), (10, 0), (10, 10), (6, 10), (6, 4), (4, 4), (4, 10), (0, 10)])
    assert point_in_polygon((0, 5), concave.vertices)
    assert point_in_polygon((2, 8), concave.vertices)
    assert not point_in_polygon((5, 8), concave.vertices)
    interacts, fully_inside = segment_room_relation((2, 8), (8, 8), concave)
    assert interacts and not fully_inside


def test_load_rooms_validation_and_configurable_extra():
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        data_dir = Path(tmp)
        (data_dir / "房间范围.csv").write_text(
            "房间名称,楼层,顶点序号,X,Y,图层,句柄\n"
            "A,1F,1,0,0,CABLE_ROOM_1F,H1\nA,1F,2,10,0,CABLE_ROOM_1F,H1\nA,1F,3,10,10,CABLE_ROOM_1F,H1\n"
            "坏房间,2F,1,0,0,CABLE_ROOM_2F,H2\n坏房间,2F,2,1,1,CABLE_ROOM_2F,H2\n",
            encoding="utf-8-sig",
        )
        (data_dir / "参数.csv").write_text("参数,值,备注\n跨房间修正,5,测试\n", encoding="utf-8-sig")
        rooms, warnings = load_rooms(data_dir)
        assert len(rooms) == 1 and rooms[0].name == "A"
        assert any("坏房间" in warning for warning in warnings)
        assert load_params(data_dir)["跨房间修正"] == 5.0


def test_app_exports_only_complete_wizard_lsp():
    import tempfile

    from cable_stat_app import export_lsp

    with tempfile.TemporaryDirectory() as tmp:
        (Path(tmp) / "cad_export_cable_route.lsp").write_text("旧文件", encoding="utf-8")
        exported = export_lsp(Path(tmp))
        assert exported == ["cad_cable_wizard.lsp"]
        assert (Path(tmp) / "cad_cable_wizard.lsp").is_file()
        assert not (Path(tmp) / "cad_export_cable_route.lsp").exists()


def test_straight_path_plus_7():
    """直线路径：基础长度 + 不拐弯修正。"""
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {
        "柜A": make_cabinet("柜A", 1000, 500),
        "柜B": make_cabinet("柜B", 9000, 500),
    }
    graph, nodes, issues = build_graph(segments, cabinets, [], ["柜A", "柜B"], PARAMS)
    assert not [i for i in issues if not i.startswith("提示")], issues
    details, _ = calculate_rows(make_rows([("柜A", "柜B")]), graph, nodes, cabinets, PARAMS)
    row = details[0]
    assert row["状态"] == "OK"
    # 沿线 8m + 两端接入各 0.5m = 9m，+7 = 16.0
    assert row["自动统计"] == 16.0, row
    assert row["向上取整"] == 16, row
    assert "路径不拐弯" in row["规则"]


def test_cabinet_names_ignore_spaces():
    """柜名匹配忽略 CAD/Excel 名称里的空格。"""
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cad_name = normalize_cabinet_name("10kV #1 主变进线柜")
    excel_name = normalize_cabinet_name("10kV#1主变进线柜")
    cabinets = {
        cad_name: make_cabinet(cad_name, 1000, 0),
        "柜B": make_cabinet("柜B", 9000, 0),
    }
    graph, nodes, issues = build_graph(segments, cabinets, [], [excel_name, "柜B"], PARAMS)
    assert not [i for i in issues if not i.startswith("提示")], issues
    details, _ = calculate_rows(make_rows([(excel_name, "柜B")]), graph, nodes, cabinets, PARAMS)
    assert details[0]["状态"] == "OK", details[0]


def test_bent_path_plus_7_then_times_1_2():
    """L形路径判定拐弯，按（CAD 测量距离 +7）×1.2。"""
    segments = [make_segment(1, 0, 0, 10000, 0), make_segment(2, 10000, 0, 10000, 10000)]
    cabinets = {
        "柜A": make_cabinet("柜A", 1000, 0),
        "柜B": make_cabinet("柜B", 10000, 9000),
    }
    graph, nodes, _ = build_graph(segments, cabinets, [], ["柜A", "柜B"], PARAMS)
    details, _ = calculate_rows(make_rows([("柜A", "柜B")]), graph, nodes, cabinets, PARAMS)
    row = details[0]
    assert row["状态"] == "OK"
    assert "拐弯" in row["规则"]
    assert "×1.2" in row["规则"]
    assert row["自动统计"] == 30.0, row  # (9m + 9m + 7) × 1.2


def test_t_junction_connects():
    """端点落在另一条线段中间时自动连通（T型连接）。"""
    # 主干线 + 一条端点搭在主干中间的支线，支线另一端接柜B
    segments = [make_segment(1, 0, 0, 20000, 0), make_segment(2, 10000, 2, 10000, 8000)]
    cabinets = {
        "柜A": make_cabinet("柜A", 1000, 0),
        "柜B": make_cabinet("柜B", 10000, 7500),
    }
    graph, nodes, issues = build_graph(segments, cabinets, [], ["柜A", "柜B"], PARAMS)
    assert any("T 型连接" in i for i in issues), issues
    details, _ = calculate_rows(make_rows([("柜A", "柜B")]), graph, nodes, cabinets, PARAMS)
    assert details[0]["状态"] == "OK", details[0]


def test_cross_junction_connects():
    """两条线段在内部十字交叉（交点不在端点），自动打通交点连通。"""
    # 水平线: (0, 5000) -> (10000, 5000)
    # 垂直线: (5000, 0) -> (5000, 10000)
    # 十字相交于 (5000, 5000)
    segments = [
        make_segment(1, 0, 5000, 10000, 5000),
        make_segment(2, 5000, 0, 5000, 10000),
    ]
    cabinets = {
        "柜A": make_cabinet("柜A", 1000, 5000),
        "柜B": make_cabinet("柜B", 5000, 9000),
    }
    graph, nodes, issues = build_graph(segments, cabinets, [], ["柜A", "柜B"], PARAMS)
    assert any("十字/交叉点" in i for i in issues), issues
    details, _ = calculate_rows(make_rows([("柜A", "柜B")]), graph, nodes, cabinets, PARAMS)
    row = details[0]
    assert row["状态"] == "OK", row
    # 柜A到交点 4m，交点到柜B 4m，总路径 8m，拐弯：(8 + 7) * 1.2 = 18.0
    assert row["基础路径长度"] == 8.0, row
    assert row["自动统计"] == 18.0, row
    assert "拐弯" in row["规则"]


def test_shaft_cross_floor():
    """跨楼层经竖井：路径含竖井高度且判为拐弯。"""
    segments = [make_segment(1, 0, 0, 10000, 0, floor="1F"), make_segment(2, 0, 0, 10000, 0, floor="2F")]
    shafts = [make_shaft("ZJ1_1F", 5000, 0, "1F"), make_shaft("ZJ1_2F", 5000, 0, "2F")]
    cabinets = {
        "柜A": make_cabinet("柜A", 1000, 0, floor="1F"),
        "柜B": make_cabinet("柜B", 9000, 0, floor="2F"),
    }
    graph, nodes, _ = build_graph(segments, cabinets, shafts, ["柜A", "柜B"], PARAMS)
    details, _ = calculate_rows(make_rows([("柜A", "柜B")]), graph, nodes, cabinets, PARAMS)
    row = details[0]
    assert row["状态"] == "OK", row
    # 1F 走 4m + 竖井 4.5m + 2F 走 4m，先 +7 再按 1.2 倍
    assert row["自动统计"] == 23.4, row
    assert row["向上取整"] == 24, row
    assert "竖井" in row["路径说明"]


def test_rule_override():
    """强制规则表覆盖自动判断的修正量。"""
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 0), "柜B": make_cabinet("柜B", 9000, 0)}
    graph, nodes, _ = build_graph(segments, cabinets, [], ["柜A", "柜B"], PARAMS)
    overrides = {("柜B", "柜A"): {"extra": 3.0, "note": "现场实测"}}  # 反向也应命中
    details, _ = calculate_rows(make_rows([("柜A", "柜B")]), graph, nodes, cabinets, PARAMS, overrides)
    row = details[0]
    assert row["修正量"] == 3.0, row
    assert "强制规则" in row["规则"] and "现场实测" in row["规则"]
    assert row["自动统计"] == 11.0, row  # 8m + 3


def test_same_start_end_rejected():
    """起点与终点相同的行标记为需人工确认。"""
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 0)}
    graph, nodes, _ = build_graph(segments, cabinets, [], ["柜A"], PARAMS)
    details, _ = calculate_rows(make_rows([("柜A", "柜A")]), graph, nodes, cabinets, PARAMS)
    assert "起点与终点相同" in details[0]["状态"]


def test_alias_loading_and_apply(tmp_path=None):
    """别名表加载与应用：清册名映射到CAD柜子，目标缺失给告警。"""
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        data_dir = Path(tmp)
        (data_dir / "柜名别名.csv").write_text(
            "清册名称,CAD名称,备注\n10kV#1主变进线柜,10kV一号主变进线柜,CAD写法不同\n不存在柜,没有这个柜,\n",
            encoding="utf-8-sig",
        )
        aliases, warnings = load_aliases(data_dir)
        assert aliases == {"10kV#1主变进线柜": "10kV一号主变进线柜", "不存在柜": "没有这个柜"}
        assert not warnings
        cabinets = {"10kV一号主变进线柜": make_cabinet("10kV一号主变进线柜", 0, 0)}
        apply_warnings = apply_aliases(cabinets, aliases)
        assert cabinets["10kV#1主变进线柜"] is cabinets["10kV一号主变进线柜"]
        assert any("不存在" in w for w in apply_warnings)


def test_rule_overrides_loading():
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        data_dir = Path(tmp)
        (data_dir / "强制规则.csv").write_text(
            "起点,终点,修正量,备注\n柜A,柜B,7,走直线\n柜A,柜B,9,重复行\n柜C,,5,缺终点\n",
            encoding="utf-8-sig",
        )
        overrides, warnings = load_rule_overrides(data_dir)
        assert overrides[("柜A", "柜B")]["extra"] == 7.0
        assert len(overrides) == 1
        assert len(warnings) == 2  # 重复 + 无效


def test_shaft_helpers():
    assert normalize_shaft_base_id("ZJ1_1F") == "ZJ1"
    assert normalize_shaft_base_id("ZJ1-2F") == "ZJ1"
    assert normalize_shaft_base_id("竖井甲二层") == "竖井甲"
    assert shaft_suffix_floor("ZJ1_1F") == "1F"
    assert shaft_suffix_floor("ZJ2_2F") == "2F"
    assert shaft_suffix_floor("ZJ3") == ""


def test_floor_elevations_for_3d():
    """三维视图楼层标高：层间距优先取跨层配对竖井 max(下层高度, 上层高度)。"""
    tall = make_shaft("ZJ1_1F", 0, 0, "1F")
    tall.height_m = 6.0
    shafts = [tall, make_shaft("ZJ1_2F", 0, 0, "2F")]
    assert compute_floor_elevations(["1F", "2F"], shafts, PARAMS) == {"1F": 0.0, "2F": 6000.0}
    # 高度标注在上层竖井时同样生效（与 build_graph 竖直边 max(h_a, h_b) 语义一致）
    tall_upper = make_shaft("ZJ1_2F", 0, 0, "2F")
    tall_upper.height_m = 6.0
    shafts_upper = [make_shaft("ZJ1_1F", 0, 0, "1F"), tall_upper]
    assert compute_floor_elevations(["1F", "2F"], shafts_upper, PARAMS) == {"1F": 0.0, "2F": 6000.0}
    # 无竖井时回退到参数"竖井高度"
    assert compute_floor_elevations(["1F", "2F"], [], PARAMS) == {"1F": 0.0, "2F": 4500.0}
    assert compute_floor_elevations([], [], PARAMS) == {}


def test_floor_offsets_for_3d():
    """三维视图楼层对齐：CAD 并排绘制的楼层平面用同名竖井平移对齐到首层。"""
    shafts = [
        make_shaft("ZJ1_1F", 1000, 2000, "1F"),
        make_shaft("ZJ2_1F", 3000, 2000, "1F"),
        make_shaft("ZJ1_2F", 61000, 2100, "2F"),
        make_shaft("ZJ2_2F", 63000, 2100, "2F"),
    ]
    offsets, warnings = compute_floor_offsets(["1F", "2F"], shafts)
    assert offsets["1F"] == {"dx": 0.0, "dy": 0.0}
    assert offsets["2F"] == {"dx": -60000.0, "dy": -100.0}
    assert not warnings
    # 无同名竖井可对齐时保持零偏移并告警
    offsets2, warnings2 = compute_floor_offsets(["1F", "2F"], shafts[:2])
    assert offsets2["2F"] == {"dx": 0.0, "dy": 0.0}
    assert any("2F" in w and "无法对齐" in w for w in warnings2)


def test_floor_offsets_chained_3_floors():
    """三层链式对齐：3F 只与 2F 有同名竖井，偏移叠加 offsets[2F]。"""
    shafts = [
        make_shaft("ZJA_1F", 1000, 2000, "1F"),
        make_shaft("ZJA_2F", 61000, 2100, "2F"),
        make_shaft("ZJB_2F", 65000, 2100, "2F"),
        make_shaft("ZJB_3F", 125000, 2200, "3F"),
    ]
    offsets, warnings = compute_floor_offsets(["1F", "2F", "3F"], shafts)
    assert offsets["2F"] == {"dx": -60000.0, "dy": -100.0}
    assert offsets["3F"] == {"dx": -120000.0, "dy": -200.0}
    assert not warnings
    # base_id 归一化要能剥掉 3F 后缀
    assert normalize_shaft_base_id("ZJB_3F") == "ZJB"
    assert normalize_shaft_base_id("ZJ5-10F") == "ZJ5"
    assert normalize_shaft_base_id("竖井甲三层") == "竖井甲"


def test_straight_collinear_segments_not_bent():
    """两段共线线段不应判为拐弯。"""
    segments = [make_segment(1, 0, 0, 5000, 0), make_segment(2, 5000, 0, 10000, 0)]
    graph = RouteGraph(segments, PARAMS)
    graph.finalize_route_edges()
    start = graph.segment_points[1][0][1]
    end = graph.segment_points[2][1][1]
    _, _, edge_metas = graph.shortest_path(start, end)
    assert not is_route_bent(graph, edge_metas)


def test_disconnected_network_reported():
    """路径网断成多块时，问题清单要列出断块明细。"""
    segments = [make_segment(1, 0, 0, 5000, 0), make_segment(2, 50000, 0, 55000, 0)]
    cabinets = {
        "柜A": make_cabinet("柜A", 1000, 500),
        "柜B": make_cabinet("柜B", 54000, 500),
    }
    graph, nodes, issues = build_graph(segments, cabinets, [], ["柜A", "柜B"], PARAMS)
    assert any("互不连通" in i for i in issues), issues
    assert any("柜A" in i for i in issues if i.strip().startswith("第")), issues
    details, _ = calculate_rows(make_rows([("柜A", "柜B")]), graph, nodes, cabinets, PARAMS)
    assert "不连通" in details[0]["状态"], details[0]


def test_unattached_cabinet_message():
    """柜子有坐标但离路径太远时，状态要说明是未接入而不是缺坐标。"""
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {
        "柜A": make_cabinet("柜A", 1000, 500),
        "远柜": make_cabinet("远柜", 1000, 999999),
    }
    graph, nodes, issues = build_graph(segments, cabinets, [], ["柜A", "远柜"], PARAMS)
    assert any("接入失败" in i for i in issues), issues
    details, _ = calculate_rows(make_rows([("柜A", "远柜")]), graph, nodes, cabinets, PARAMS)
    assert "未接入路径网" in details[0]["状态"], details[0]


def test_param_sanity_hints():
    """吸附容差过大、接入距离不够时，问题清单给出带建议值的参数提示。"""
    params = dict(PARAMS)
    params["吸附容差"] = 6000.0
    params["最大接入距离"] = 100.0
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {
        "柜A": make_cabinet("柜A", 1000, 500),
        "柜B": make_cabinet("柜B", 5000, 3000),
    }
    graph, nodes, issues = build_graph(segments, cabinets, [], ["柜A", "柜B"], params)
    assert any(i.startswith("参数提示") and "吸附容差" in i for i in issues), issues
    assert any(i.startswith("参数提示") and "最大接入距离" in i and "3901" in i for i in issues), issues


def test_workbook_headers_tolerant():
    """表头带换行/空格要能识别；缺「自动统计」「向上取整」列时自动在表尾补建。"""
    import tempfile

    import openpyxl

    from calculate_cable_lengths import load_workbook_rows

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["分类", "电缆编号", "起 点", "终点", "电缆\n长度"])
    ws.append(["", "1U-101", "A柜", "B柜", 8])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "自动统计.xlsx"
        wb.save(path)
        wb2, ws2, headers, rows, names = load_workbook_rows(path)
        assert headers["电缆长度"] == 5
        assert headers["起点"] == 3
        # 自动补建在最后一个表头之后，「向上取整」紧跟「自动统计」右边
        assert headers["自动统计"] == 6
        assert ws2.cell(1, 6).value == "自动统计"
        assert headers["向上取整"] == 7
        assert ws2.cell(1, 7).value == "向上取整"
        assert rows[0]["起点"] == "A柜"
        assert rows[0]["终点"] == "B柜"


def test_workbook_ceil_column_inserted_next_to_auto():
    """「自动统计」右边已有其他列时，「向上取整」插在它右边且原数据整体右移。"""
    import tempfile

    import openpyxl

    from calculate_cable_lengths import load_workbook_rows

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["电缆编号", "起点", "终点", "电缆长度", "自动统计", "备注"])
    ws.append(["1U-101", "A柜", "B柜", 8, 9.5, "留观"])
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "自动统计.xlsx"
        wb.save(path)
        wb2, ws2, headers, rows, names = load_workbook_rows(path)
        assert headers["自动统计"] == 5
        assert headers["向上取整"] == 6
        assert ws2.cell(1, 6).value == "向上取整"
        # 原「备注」列被顺移到第 7 列，数据不丢
        assert headers["备注"] == 7
        assert ws2.cell(2, 7).value == "留观"
        assert rows[0]["起点"] == "A柜"


def test_workbook_formula_manual_length_uses_cached_value():
    """「电缆长度」是公式时取 Excel 上次保存的计算值，差值才能算出来。"""
    import re
    import tempfile
    import zipfile

    import openpyxl

    from calculate_cable_lengths import load_workbook_rows

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["电缆编号", "起点", "终点", "电缆长度", "备注"])
    ws.append(["1U-101", "A柜", "B柜", "=5+7.5", "x"])
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "raw.xlsx"
        path = Path(tmp) / "自动统计.xlsx"
        wb.save(raw)
        # openpyxl 不算公式，这里模拟 Excel 保存后留下的缓存值
        with zipfile.ZipFile(raw) as src, zipfile.ZipFile(path, "w") as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    text = re.sub(r"<f>5\+7\.5</f>(<v\s*/>|<v></v>)?", "<f>5+7.5</f><v>12.5</v>", data.decode("utf-8"))
                    data = text.encode("utf-8")
                dst.writestr(item, data)
        wb2, ws2, headers, rows, _ = load_workbook_rows(path)
        assert rows[0]["手工长度"] == 12.5, rows
        # 插入「向上取整」列后公式单元格本身保持不变
        assert ws2.cell(2, headers["电缆长度"]).value == "=5+7.5"


def test_find_workbook_requires_cable_no_header():
    import tempfile

    import openpyxl

    from run_auto_stat import find_workbook

    with tempfile.TemporaryDirectory() as tmp:
        wb = openpyxl.Workbook()
        wb.active.append(["起点", "终点", "电缆长度"])
        wb.save(Path(tmp) / "缺编号.xlsx")
        try:
            find_workbook(Path(tmp))
        except FileNotFoundError:
            pass
        else:
            raise AssertionError("缺「电缆编号」列的表不应被当成清册")
        wb = openpyxl.Workbook()
        wb.active.append(["电缆编号", "起点", "终点", "电缆\n长度"])
        wb.save(Path(tmp) / "清册.xlsx")
        assert find_workbook(Path(tmp)).name == "清册.xlsx"


def test_suspicious_rooms_reported():
    """房间里既没柜子也没路径（误放在房间图层的图框/表格）、以及未命名房间都要提示。"""
    from calculate_cable_lengths import check_suspicious_rooms

    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 500)}
    rooms = [
        make_room("配电室", [(0, 0), (2000, 0), (2000, 2000), (0, 2000)]),
        make_room("走廊", [(4000, -500), (6000, -500), (6000, 500), (4000, 500)]),
        make_room("图框", [(600000, 0), (607000, 0), (607000, 6000), (600000, 6000)]),
        make_room("房间_1F130", [(4000, -500), (6000, -500), (6000, 500), (4000, 500)], floor="2F"),
    ]
    warnings = check_suspicious_rooms(rooms, cabinets, segments)
    assert not any("配电室" in w or "走廊" in w for w in warnings), warnings
    assert any("图框" in w and "既没有柜子也没有路径" in w for w in warnings), warnings
    assert any("房间_1F130" in w and "没有名称" in w for w in warnings), warnings


def test_locked_output_workbook_saved_under_new_name():
    import tempfile

    import openpyxl

    from calculate_cable_lengths import write_results_to_workbook

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["电缆编号", "起点", "终点", "电缆长度", "自动统计", "向上取整"])
    real_save = wb.save
    attempts: list[Path] = []

    def save(path):
        attempts.append(Path(path))
        if len(attempts) == 1:
            raise PermissionError("locked")
        real_save(path)

    wb.save = save
    with tempfile.TemporaryDirectory() as tmp:
        target = Path(tmp) / "outputs" / "自动统计_计算结果.xlsx"
        saved = write_results_to_workbook(
            target, target, wb, ws, {"自动统计": 5, "向上取整": 6}, [], dict(PARAMS), [], []
        )
        assert saved != target and saved.exists(), (saved, attempts)
        assert saved.name.startswith("自动统计_计算结果_") and saved.suffix == ".xlsx"


def test_visual_edges_reference_node_ids():
    """可视化数据：边表去重、只存节点编号；电缆 path 存边序号（~k 为反向），首尾衔接成连续路径。"""
    from calculate_cable_lengths import build_visualization_data

    segments = [make_segment(1, 0, 0, 10000, 0), make_segment(2, 10000, 0, 10000, 5000)]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 100), "柜B": make_cabinet("柜B", 10100, 4000)}
    graph, nodes, _ = build_graph(segments, cabinets, [], list(cabinets), PARAMS)
    _, visuals = calculate_rows(make_rows([("柜A", "柜B"), ("柜B", "柜A")]), graph, nodes, cabinets, PARAMS)
    data = build_visualization_data(Path("a.xlsx"), Path("b.xlsx"), PARAMS, graph, segments, cabinets, [], [], visuals, [])
    assert data["version"] == 4
    node_ids = {node["id"] for node in data["nodes"]}
    edges = data["edges"]
    assert edges and all(e["from"] in node_ids and e["to"] in node_ids for e in edges)
    forward, backward = data["cables"][0]["path"], data["cables"][1]["path"]
    # 反向电缆复用同一批边，只是方向相反
    assert len(edges) == len(forward) and sorted(~k for k in backward) == sorted(forward)

    def walk(path):
        steps = [(edges[k]["from"], edges[k]["to"]) if k >= 0 else (edges[~k]["to"], edges[~k]["from"]) for k in path]
        assert all(a[1] == b[0] for a, b in zip(steps, steps[1:])), steps
        return steps

    first, second = walk(forward), walk(backward)
    assert first[0][0] == second[-1][1] and first[-1][1] == second[0][0]
    assert "path_metas" not in data["cables"][0] and "path_edges" not in data["cables"][0]


def test_floor_normalization_variants():
    from calculate_cable_lengths import floor_sort_key, infer_floor_from_layer, normalize_floor

    cases = {
        "1": "1F", "F2": "2F", "3f": "3F", "一层": "1F", "二楼": "2F", "十一层": "11F", "二十层": "20F",
        "12层": "12F", "B1": "B1F", "b2f": "B2F", "-1F": "B1F", "地下一层": "B1F", "负二层": "B2F", "": "",
    }
    for raw, expected in cases.items():
        assert normalize_floor(raw) == expected, (raw, normalize_floor(raw))
    assert infer_floor_from_layer("CABLE_ROUTE_12F") == "12F"
    assert infer_floor_from_layer("CABLE_ROUTE_B1F") == "B1F"
    assert infer_floor_from_layer("电缆沟三层") == "3F"
    assert infer_floor_from_layer("0") == ""
    assert sorted(["2F", "B1F", "1F", "B2F", "10F"], key=floor_sort_key) == ["B2F", "B1F", "1F", "2F", "10F"]
    assert normalize_shaft_base_id("ZJ1_B1F") == "ZJ1"
    assert shaft_suffix_floor("ZJ1_B1F") == "B1F"
    assert shaft_suffix_floor("ZJ1地下一层") == "B1F"


def test_fullwidth_cabinet_names_match():
    assert normalize_cabinet_name("１＃主变 保护柜") == "1#主变保护柜"
    assert normalize_cabinet_name("10kV Ⅰ段") == "10kVⅠ段"


def test_round_half_up_and_manual_length_parsing():
    from cable_stat.text import parse_length, round_half_up

    assert round_half_up(7.25, 1) == 7.3
    assert round_half_up(23.449999999999996, 1) == 23.5
    assert round_half_up(-2.25, 1) == -2.3
    assert round_half_up(16.04, 1) == 16.0
    assert parse_length("12") == 12.0 and parse_length("12.5m") == 12.5 and parse_length("1,200米") == 1200.0
    assert parse_length("约10") is None and parse_length(None) is None and parse_length(True) is None


def test_bend_angle_parameter():
    """轻微歪斜（约 3°）在默认 5° 判定角下不算拐弯，判定角调到 2° 就算。"""
    segments = [make_segment(1, 0, 0, 5000, 0), make_segment(2, 5000, 0, 10000, 262)]
    cabinets = {"柜A": make_cabinet("柜A", 500, 0), "柜B": make_cabinet("柜B", 9500, 236)}
    for angle, expect_bent in ((5.0, False), (2.0, True)):
        params = dict(PARAMS, 拐弯判定角度=angle)
        graph, nodes, _ = build_graph(segments, cabinets, [], list(cabinets), params)
        details, _ = calculate_rows(make_rows([("柜A", "柜B")]), graph, nodes, cabinets, params)
        assert ("拐弯" in details[0]["规则"] and "不拐弯" not in details[0]["规则"]) == expect_bent, (angle, details[0])


def test_override_negative_extra_label_and_empty_endpoints():
    segments = [make_segment(1, 0, 0, 10000, 0)]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 0), "柜B": make_cabinet("柜B", 9000, 0)}
    graph, nodes, _ = build_graph(segments, cabinets, [], ["柜A", "柜B"], PARAMS)
    overrides = {("柜A", "柜B"): {"extra": -2.0, "note": ""}}
    rows = make_rows([("柜A", "柜B"), ("", "柜B"), ("柜A", "")])
    details, _ = calculate_rows(rows, graph, nodes, cabinets, PARAMS, overrides)
    assert details[0]["规则"] == "强制规则 -2" and details[0]["自动统计"] == 6.0, details[0]
    assert details[1]["状态"] == "清册起点为空" and details[2]["状态"] == "清册终点为空"


def test_params_validation_and_save_roundtrip():
    import tempfile

    from cable_stat.params import load_params_with_warnings, save_params

    with tempfile.TemporaryDirectory() as tmp:
        data_dir = Path(tmp)
        (data_dir / "参数.csv").write_text(
            "参数,值,备注\n拐弯倍率,abc,\n吸附容差,-1,\n吸附距离,5,写错名字\n竖井高度,6,自定义备注\n",
            encoding="utf-8-sig",
        )
        params, warnings = load_params_with_warnings(data_dir)
        assert params["拐弯倍率"] == 1.2 and params["吸附容差"] == 0.05 and params["竖井高度"] == 6.0
        assert any("拐弯倍率" in w and "不是数字" in w for w in warnings)
        assert any("吸附容差" in w and "不能小于" in w for w in warnings)
        assert any("未知参数" in w and "吸附距离" in w for w in warnings)
        save_params(data_dir, {"竖井高度": 5.5, "跨房间修正": 2})
        text = (data_dir / "参数.csv").read_text(encoding="utf-8-sig")
        assert "竖井高度,5.5,自定义备注" in text and "跨房间修正,2," in text
        assert load_params(data_dir)["竖井高度"] == 5.5
        (data_dir / "参数.csv").write_text("参数,值\nCAD每米单位,0\n", encoding="utf-8-sig")
        try:
            load_params(data_dir)
        except ValueError:
            pass
        else:
            raise AssertionError("CAD每米单位=0 应直接报错")


def test_workbook_title_rows_merged_cells_and_model_column():
    """表头在第 3 行；起点柜纵向合并；横向合并的分组标题行不能当成电缆；型号+规格两列拼接；字符串长度能算差值。"""
    import tempfile

    import openpyxl

    from calculate_cable_lengths import load_workbook_rows

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["XX变电站电缆清册"])
    ws.append(["图号：D-01"])
    ws.append(["电缆编号", "型号", "规格", "起点", "终点", "电缆长度"])
    ws.append(["一、控制电缆"])
    ws.merge_cells("A4:F4")
    ws.append(["C1", "KVVP", "4×1.5", "A柜", "B柜", "12米"])
    ws.append(["C2", "KVVP", "4×1.5", None, "C柜", 8])
    ws.merge_cells("D5:D6")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "清册.xlsx"
        wb.save(path)
        _, ws2, headers, rows, names = load_workbook_rows(path)
        assert headers["起点"] == 4 and headers["自动统计"] == 7 and headers["向上取整"] == 8
        assert ws2.cell(3, 7).value == "自动统计"
        assert [r["电缆编号"] for r in rows] == ["C1", "C2"], rows
        assert [r["起点"] for r in rows] == ["A柜", "A柜"]
        assert rows[0]["手工长度"] == 12.0 and rows[0]["型号"] == "KVVP 4×1.5"
        assert names == ["A柜", "B柜", "C柜"]


def test_template_params_match_registry():
    """项目模板的 参数.csv 与内置参数登记表（名称、默认值）保持一致。"""
    from cable_stat.csvio import read_csv_dicts
    from cable_stat.params import PARAM_SPECS

    rows = read_csv_dicts(Path(__file__).resolve().parent / "项目模板" / "data" / "参数.csv")
    assert [row["参数"] for row in rows] == [spec.name for spec in PARAM_SPECS]
    for row, spec in zip(rows, PARAM_SPECS):
        assert float(row["值"]) == spec.default, (spec.name, row["值"])
        assert row["备注"] == spec.note, spec.name


def test_insert_ceil_column_keeps_formulas_and_merges():
    import tempfile

    import openpyxl

    from calculate_cable_lengths import load_workbook_rows

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["电缆编号", "起点", "终点", "电缆长度", "自动统计", "备注", "校核"])
    ws.append(["C1", "A柜", "B柜", 8, 9.5, "x", "=E2-D2+F2"])
    ws.append(["C2", "A柜", "C柜", 5, 6.5, "合并备注"])
    ws.merge_cells("F3:G3")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "自动统计.xlsx"
        wb.save(path)
        _, ws2, headers, _, _ = load_workbook_rows(path)
        assert headers["向上取整"] == 6 and headers["校核"] == 8
        # 公式里指向右移列的引用跟着右移，左侧列不动
        assert ws2.cell(2, 8).value == "=E2-D2+G2", ws2.cell(2, 8).value
        assert "G3:H3" in {str(r) for r in ws2.merged_cells.ranges}
        assert ws2.cell(3, 7).value == "合并备注"


def test_case_insensitive_cabinet_match_and_duplicate_ids():
    from cable_stat.loaders import apply_case_insensitive_matches
    from cable_stat.pipeline import duplicate_cable_ids

    cabinets = {"10kV进线柜": make_cabinet("10kV进线柜", 0, 0)}
    notes = apply_case_insensitive_matches(cabinets, ["10KV进线柜", "不存在"])
    assert cabinets["10KV进线柜"] is cabinets["10kV进线柜"] and len(notes) == 1
    rows = make_rows([("A", "B"), ("A", "C"), ("B", "C")])
    rows[2]["电缆编号"] = "C1"
    warnings = duplicate_cable_ids(rows)
    assert len(warnings) == 1 and "C1" in warnings[0] and "2、4" in warnings[0]


def test_issue_levels_and_summary():
    from cable_stat.report import build_summary, classify_issues

    levels = classify_issues([
        "已计算 3 条，未计算 1 条",
        "共有 1 条电缆未能计算，详见“统计明细”状态列",
        "参数提示：吸附容差过大",
        "固定路径网络分成 2 块互不连通（…），断块明细：",
        "  第1块（楼层 1F，3个节点）挂载：柜A",
        "柜子名称重复：柜A，已采用第一条坐标",
    ])
    assert [level for level, _ in levels] == ["信息", "错误", "提示", "错误", "错误", "警告"]
    detail = lambda model, auto, status="OK", manual=None: {  # noqa: E731
        "型号": model, "自动统计": auto, "向上取整": int(-(-auto // 1)) if status == "OK" else "", "状态": status,
        "规则": "路径不拐弯 +7", "手工长度": manual, "差值": (round(auto - manual, 1) if manual is not None else ""), "跨房间": "否",
    }
    rows = [detail("A型", 10.2, manual=4), detail("A型", 5.0), detail("B型", 3.1), detail("B型", "", status="缺少起点柜坐标：X")]
    summary = build_summary(rows, dict(PARAMS, 差值提醒=5))
    assert summary["ok"] == 3 and summary["failed"] == 1 and summary["ceil_sum"] == 20 and summary["diff_alerts"] == 1
    by_model = {item["型号"]: item for item in summary["by_model"]}
    assert by_model["A型"]["向上取整合计"] == 16 and by_model["B型"]["条数"] == 2 and by_model["B型"]["已计算"] == 1


def test_project_scan_and_alias_save():
    import tempfile

    import openpyxl

    from cable_stat.project import load_alias_rows, save_aliases, scan_project

    with tempfile.TemporaryDirectory() as tmp:
        project = Path(tmp)
        (project / "data").mkdir()
        wb = openpyxl.Workbook()
        wb.active.append(["电缆编号", "起点", "终点", "电缆长度"])
        wb.save(project / "自动统计.xlsx")
        (project / "data" / "柜子坐标.csv").write_text("柜子名称,X,Y\nA,0,0\nB,1,1\n", encoding="utf-8-sig")
        (project / "data" / "柜名别名.csv").write_text("清册名称,CAD名称,备注\n旧名,A,保留\n改名,B,旧\n", encoding="utf-8-sig")
        status = scan_project(project)
        assert status.workbook and status.workbook.name == "自动统计.xlsx"
        assert status.missing_required == ["路径线段.csv"]
        assert {f.name: f.rows for f in status.files}["柜子坐标.csv"] == 2
        assert not status.has_result
        assert save_aliases(project / "data", {"改名": "A", "新名": "B"}) == 2
        rows = {r["清册名称"]: r["CAD名称"] for r in load_alias_rows(project / "data")}
        assert rows == {"旧名": "A", "改名": "A", "新名": "B"}


def test_grid_index_matches_brute_force():
    """网格索引的最近线段/吸附结果与逐个比较完全一致（含远离路径网的点、超长斜线）。"""
    import math
    import random

    from cable_stat.geometry import PointGrid, SegmentGrid, project_to_segment

    rnd = random.Random(7)
    segments = [
        make_segment(i, rnd.uniform(0, 50000), rnd.uniform(0, 50000), rnd.uniform(0, 50000), rnd.uniform(0, 50000))
        for i in range(1, 80)
    ]
    segments.append(make_segment(80, -1e7, -1e7, 1e7, 1e7))
    grid = SegmentGrid(segments, pad=10)
    for _ in range(300):
        x, y = rnd.uniform(-20000, 70000), rnd.uniform(-20000, 70000)
        brute = min(segments, key=lambda s: (project_to_segment(x, y, s)[3], s.idx))
        found = grid.nearest(x, y)
        assert found[2].idx == brute.idx, (x, y)
    points = PointGrid(5.0)
    coords = [(rnd.uniform(0, 100), rnd.uniform(0, 100)) for _ in range(200)]
    for i, (px, py) in enumerate(coords):
        points.add(i, px, py)
    for _ in range(300):
        x, y = rnd.uniform(0, 100), rnd.uniform(0, 100)
        within = [(math.hypot(x - px, y - py), i) for i, (px, py) in enumerate(coords) if math.hypot(x - px, y - py) <= 5.0]
        assert points.nearest_within(x, y, 5.0) == (min(within)[1] if within else None)


def test_route_gap_detection():
    """线头差一点没搭上：导致断网的报“错误”且只报一次；网内缩头（会绕路）报“疑似断点”；参数为0不检查。"""
    from cable_stat.report import issue_level

    params = dict(PARAMS, 断点提示距离=0.5)
    split = [make_segment(1, 0, 0, 10000, 0), make_segment(2, 10030, 0, 20000, 0)]
    cabinets = {"柜A": make_cabinet("柜A", 1000, 0), "柜B": make_cabinet("柜B", 19000, 0)}
    graph, _, issues = build_graph(split, cabinets, [], list(cabinets), params)
    gap_issues = [i for i in issues if i.startswith("疑似断点")]
    assert len(graph.gaps) == 1 and len(gap_issues) == 1, issues
    assert graph.gaps[0]["bridging"] and abs(graph.gaps[0]["dist"] - 30) < 1e-6
    assert issue_level(gap_issues[0]) == "错误" and "只差 30.0" in gap_issues[0]

    loop = [
        make_segment(1, 0, 0, 10000, 0),
        make_segment(2, 2000, 0, 2000, 3000),
        make_segment(3, 0, 3020, 10000, 3020),
        make_segment(4, 10000, 0, 10000, 3020),
    ]
    graph, _, issues = build_graph(loop, cabinets, [], [], params)
    assert [g["route_id"] for g in graph.gaps] == ["R3"] and not graph.gaps[0]["bridging"], graph.gaps
    assert abs(graph.gaps[0]["dist"] - 20) < 1e-6 and graph.gaps[0]["x"] == 2000 and graph.gaps[0]["py"] == 3020
    assert any(i.startswith("疑似断点：") and "绕远路" in i for i in issues), issues
    graph, _, issues = build_graph(loop, cabinets, [], [], dict(PARAMS, 断点提示距离=0))
    assert graph.gaps == [] and not any("疑似断点" in i for i in issues)


def _run_all() -> int:
    failed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"PASS {name}")
            except Exception as exc:
                failed += 1
                print(f"FAIL {name}: {type(exc).__name__}: {exc}")
    return failed


if __name__ == "__main__":
    sys.exit(_run_all())
