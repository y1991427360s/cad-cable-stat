"""结果输出到 Excel：原清册填值 + 汇总 / 统计明细 / 柜子清单 / 参数 / 问题清单。"""

from __future__ import annotations

import difflib
import os
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .calc import DETAIL_HEADERS, RULE_BENT, RULE_OVERRIDE, RULE_STRAIGHT
from .models import Cabinet
from .params import PARAM_SPEC_BY_NAME

CABINET_CHECK_HEADERS = ["柜子名称", "清册出现次数", "已提供坐标", "楼层", "X", "Y", "近似CAD柜名建议"]
ISSUE_HEADERS = ["级别", "说明"]

LEVEL_ERROR = "错误"
LEVEL_WARN = "警告"
LEVEL_HINT = "提示"
LEVEL_INFO = "信息"

_HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_BAD_FILL = PatternFill("solid", fgColor="F8D7DA")
_WARN_FILL = PatternFill("solid", fgColor="FFF3CD")
_HINT_FILL = PatternFill("solid", fgColor="E7F1FA")
_LEVEL_FILLS = {LEVEL_ERROR: _BAD_FILL, LEVEL_WARN: _WARN_FILL, LEVEL_HINT: _HINT_FILL}

_ERROR_PREFIXES = ("柜子接入失败", "竖井接入失败", "清册柜子未在", "固定路径网络分成", "缺少路径线段", "共有", "疑似断点（导致断网）")


def issue_level(text: str) -> str:
    if text.startswith("已计算"):
        return LEVEL_INFO
    if text.startswith(("提示", "参数提示", "三维视图")):
        return LEVEL_HINT
    if text.startswith(_ERROR_PREFIXES):
        return LEVEL_ERROR
    return LEVEL_WARN


def classify_issues(issues: list[str]) -> list[tuple[str, str]]:
    """给问题清单每条定级；缩进的续行（如断块明细“  第1块…”）沿用上一条的级别。"""
    result: list[tuple[str, str]] = []
    for text in issues:
        if text.startswith(" ") and result:
            result.append((result[-1][0], text))
        else:
            result.append((issue_level(text), text))
    return result


def display_width(value: Any) -> float:
    text = str(value)
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)


def remove_sheet_if_exists(wb: Any, name: str) -> None:
    if name in wb.sheetnames:
        del wb[name]


def style_header(ws: Any, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def fit_columns(ws: Any, max_width: int = 40) -> None:
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = 8.0
        for row in range(1, min(ws.max_row, 200) + 1):
            value = ws.cell(row, col).value
            if value is None:
                continue
            width = max(width, min(max_width, display_width(value) * 1.1 + 2))
        ws.column_dimensions[letter].width = width


def write_sheet(wb: Any, name: str, headers: list[str], rows: list[dict[str, Any]]) -> Any:
    remove_sheet_if_exists(wb, name)
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fit_columns(ws, max_width=50)
    return ws


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def diff_alert(detail: dict[str, Any], params: dict[str, float]) -> bool:
    threshold = params.get("差值提醒", 0.0)
    return threshold > 0 and _is_number(detail.get("差值")) and abs(detail["差值"]) >= threshold


def build_summary(detail_rows: list[dict[str, Any]], params: dict[str, float]) -> dict[str, Any]:
    """总体数量与长度合计、按规则分类计数、按型号汇总长度（清册有型号列时）。"""
    ok_rows = [row for row in detail_rows if row["状态"] == "OK"]
    manual_rows = [row for row in ok_rows if _is_number(row["手工长度"])]
    by_rule: dict[str, int] = defaultdict(int)
    for row in ok_rows:
        rule = str(row["规则"])
        kind = RULE_OVERRIDE if rule.startswith("强制规则") else RULE_BENT if "拐弯/跨楼层" in rule else RULE_STRAIGHT
        by_rule[kind] += 1
    by_model: dict[str, dict[str, Any]] = {}
    for row in detail_rows:
        model = row.get("型号") or ""
        item = by_model.setdefault(model, {"型号": model or "（未填型号）", "条数": 0, "已计算": 0, "自动统计合计": 0.0, "向上取整合计": 0})
        item["条数"] += 1
        if row["状态"] == "OK":
            item["已计算"] += 1
            item["自动统计合计"] += row["自动统计"]
            item["向上取整合计"] += row["向上取整"]
    models = sorted(by_model.values(), key=lambda item: (item["型号"] == "（未填型号）", item["型号"]))
    for item in models:
        item["自动统计合计"] = round(item["自动统计合计"], 1)
    return {
        "total": len(detail_rows),
        "ok": len(ok_rows),
        "failed": len(detail_rows) - len(ok_rows),
        "auto_sum": round(sum(row["自动统计"] for row in ok_rows), 1),
        "ceil_sum": sum(row["向上取整"] for row in ok_rows),
        "manual_count": len(manual_rows),
        "manual_sum": round(sum(float(row["手工长度"]) for row in manual_rows), 1),
        "auto_sum_with_manual": round(sum(row["自动统计"] for row in manual_rows), 1),
        "diff_alerts": sum(1 for row in ok_rows if diff_alert(row, params)),
        "cross_room": sum(1 for row in ok_rows if row["跨房间"] == "是"),
        "by_rule": dict(by_rule),
        "by_model": models if any(row.get("型号") for row in detail_rows) else [],
    }


def write_summary_sheet(wb: Any, summary: dict[str, Any], params: dict[str, float]) -> Any:
    remove_sheet_if_exists(wb, "汇总")
    ws = wb.create_sheet("汇总")
    bold = Font(bold=True)
    title = Font(bold=True, size=12)

    started = False

    def section(name: str) -> None:
        nonlocal started
        if started:
            ws.append([])
        started = True
        ws.append([name])
        ws.cell(ws.max_row, 1).font = title

    section("总体")
    rows = [
        ("电缆总数", summary["total"], "条"),
        ("已计算", summary["ok"], "条"),
        ("未计算", summary["failed"], "条（原因见“统计明细”状态列）"),
        ("自动统计合计", summary["auto_sum"], "米"),
        ("向上取整合计", summary["ceil_sum"], "米"),
        ("有手工长度的电缆", summary["manual_count"], "条"),
        ("其手工长度合计", summary["manual_sum"], "米"),
        ("其自动统计合计", summary["auto_sum_with_manual"], "米"),
        (f"差值超过 {params.get('差值提醒', 0):g} 米", summary["diff_alerts"], "条（统计明细里标黄）"),
        ("跨房间", summary["cross_room"], "条"),
    ]
    for row in rows:
        ws.append(list(row))
    section("按规则")
    for kind in (RULE_STRAIGHT, RULE_BENT, RULE_OVERRIDE):
        ws.append([kind, summary["by_rule"].get(kind, 0), "条"])
    if summary["by_model"]:
        section("按型号（向上取整合计可直接作为材料量参考）")
        ws.append(["型号", "条数", "已计算", "自动统计合计(米)", "向上取整合计(米)"])
        for cell in ws[ws.max_row]:
            cell.font = bold
            cell.fill = _HEADER_FILL
        for item in summary["by_model"]:
            ws.append([item["型号"], item["条数"], item["已计算"], item["自动统计合计"], item["向上取整合计"]])
    ws.column_dimensions["A"].width = 34
    for letter in "BCDE":
        ws.column_dimensions[letter].width = 16
    return ws


def write_results_to_workbook(
    workbook_path: Path,
    output_path: Path,
    wb: Any,
    ws: Any,
    headers: dict[str, int],
    detail_rows: list[dict[str, Any]],
    params: dict[str, float],
    cabinet_check_rows: list[dict[str, Any]],
    issues: list[str],
) -> Path:
    """写结果工作簿，返回实际保存路径；结果文件正被 Excel 打开时改存带时间的新文件名。"""
    auto_col = headers["自动统计"]
    ceil_col = headers["向上取整"]
    for detail in detail_rows:
        row_no = detail["行号"]
        value = detail["自动统计"]
        auto_cell = ws.cell(row_no, auto_col)
        auto_cell.value = value if _is_number(value) else None
        auto_cell.number_format = "0.0"
        ceil_value = detail["向上取整"]
        ceil_cell = ws.cell(row_no, ceil_col)
        ceil_cell.value = ceil_value if _is_number(ceil_value) else None
        ceil_cell.number_format = "0"
        # 只标记工具自己写的两列：没算出来的标红并把原因写进批注，差值过大的标黄
        if detail["状态"] != "OK":
            auto_cell.fill = _BAD_FILL
            auto_cell.comment = Comment(f"未计算：{detail['状态']}", "电缆统计")
        elif diff_alert(detail, params):
            auto_cell.fill = _WARN_FILL
            auto_cell.comment = Comment(f"与手工长度相差 {detail['差值']} 米", "电缆统计")
    for col in (auto_col, ceil_col):
        dim = ws.column_dimensions[get_column_letter(col)]
        if not dim.width or dim.width < 9:
            dim.width = 10

    summary = build_summary(detail_rows, params)
    write_summary_sheet(wb, summary, params)
    detail_ws = write_sheet(wb, "统计明细", DETAIL_HEADERS, detail_rows)
    status_col = DETAIL_HEADERS.index("状态") + 1
    diff_col = DETAIL_HEADERS.index("差值") + 1
    for offset, detail in enumerate(detail_rows, start=2):
        if detail["状态"] != "OK":
            detail_ws.cell(offset, status_col).fill = _BAD_FILL
        elif diff_alert(detail, params):
            detail_ws.cell(offset, diff_col).fill = _WARN_FILL
    write_sheet(wb, "柜子清单", CABINET_CHECK_HEADERS, cabinet_check_rows)
    write_sheet(
        wb,
        "参数",
        ["参数", "值", "说明"],
        [
            {"参数": key, "值": value, "说明": PARAM_SPEC_BY_NAME[key].note if key in PARAM_SPEC_BY_NAME else ""}
            for key, value in params.items()
        ],
    )
    classified = classify_issues(issues) or [(LEVEL_INFO, "未发现结构性问题")]
    issue_ws = write_sheet(wb, "问题清单", ISSUE_HEADERS, [{"级别": level, "说明": text} for level, text in classified])
    issue_ws.column_dimensions["B"].width = 120
    for offset, (level, _) in enumerate(classified, start=2):
        fill = _LEVEL_FILLS.get(level)
        if fill:
            issue_ws.cell(offset, 1).fill = fill
    return save_workbook(wb, output_path)


def save_workbook(wb: Any, output_path: Path) -> Path:
    """先存临时文件再替换；目标正被 Excel 占用时另存为带时间的新文件名。"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fallback = output_path.with_name(f"{output_path.stem}_{datetime.now():%H%M%S}{output_path.suffix}")
    tmp = output_path.with_name(f"~tmp_{output_path.name}")
    try:
        wb.save(tmp)
    except PermissionError:
        wb.save(fallback)
        print(f"注意：{output_path.name} 正被占用（多半在 Excel 里开着），本次结果另存为 {fallback.name}")
        return fallback
    try:
        os.replace(tmp, output_path)
        return output_path
    except PermissionError:
        os.replace(tmp, fallback)
        print(f"注意：{output_path.name} 正被占用（多半在 Excel 里开着），本次结果另存为 {fallback.name}")
        return fallback


def make_cabinet_check_rows(rows: list[dict[str, Any]], cabinets: dict[str, Cabinet]) -> list[dict[str, Any]]:
    counts: dict[str, int] = defaultdict(int)
    for row in rows:
        if row["起点"]:
            counts[row["起点"]] += 1
        if row["终点"]:
            counts[row["终点"]] += 1
    candidate_names = sorted(cabinets)
    result = []
    for name, count in sorted(counts.items(), key=lambda item: (-item[1], item[0])):
        cab = cabinets.get(name)
        suggestion = ""
        if not cab and candidate_names:
            suggestion = "；".join(suggest_cabinet_names(name, candidate_names))
        result.append(
            {
                "柜子名称": name,
                "清册出现次数": count,
                "已提供坐标": "是" if cab else "否",
                "楼层": cab.floor if cab else "",
                "X": cab.x if cab else "",
                "Y": cab.y if cab else "",
                "近似CAD柜名建议": suggestion,
            }
        )
    return result


def suggest_cabinet_names(name: str, candidates: list[str], limit: int = 3) -> list[str]:
    """近似 CAD 柜名：先比较忽略大小写的相似度，再补上互相包含的名字（如“1#主变保护柜”与“1#主变保护柜A”）。"""
    folded = {cand.casefold(): cand for cand in candidates}
    matches = [folded[m] for m in difflib.get_close_matches(name.casefold(), list(folded), n=limit, cutoff=0.55)]
    if len(matches) < limit and len(name) >= 3:
        for cand in candidates:
            if cand not in matches and (name in cand or cand in name) and min(len(cand), len(name)) >= 3:
                matches.append(cand)
                if len(matches) >= limit:
                    break
    return matches
