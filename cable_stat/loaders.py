"""输入数据读取：data\\ 里的 CSV（柜子、路径、竖井、房间、别名、强制规则）和电缆清册工作簿。"""

from __future__ import annotations

import math
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter

from .csvio import data_csv_path, read_csv_dicts
from .geometry import polygon_signed_area
from .models import Cabinet, Room, Segment, Shaft
from .text import (
    infer_floor_from_layer,
    normalize_cabinet_name,
    normalize_floor,
    normalize_header,
    normalize_text,
    parse_float,
    parse_length,
)

# 清册必需表头；「自动统计」「向上取整」列可以没有，读取时自动补建
REQUIRED_HEADERS = ("电缆编号", "起点", "终点", "电缆长度")
# 可选的电缆型号列，有的话结果里按型号汇总长度（按顺序取第一个出现的）；
# 型号和规格分两列写（ZR-KVVP | 4×1.5）时两列拼起来作为汇总键
MODEL_HEADERS = ("型号", "电缆型号", "规格型号", "型号规格", "型号及规格")
SPEC_HEADERS = ("规格",)
# 表头不一定在第一行（上面常有工程名称、图号等标题行），在前几行里找
HEADER_SCAN_ROWS = 10


def load_cabinets(data_dir: Path) -> tuple[dict[str, Cabinet], list[str]]:
    cabinets: dict[str, Cabinet] = {}
    warnings: list[str] = []
    for row in read_csv_dicts(data_csv_path(data_dir, "柜子坐标.csv", "cabinets.csv")):
        name = normalize_cabinet_name(row.get("柜子名称"))
        x = parse_float(row.get("X"))
        y = parse_float(row.get("Y"))
        if not name or x is None or y is None:
            continue
        floor = normalize_floor(row.get("楼层")) or infer_floor_from_layer(row.get("图层", ""))
        cabinet = Cabinet(
            name=name,
            floor=floor,
            x=x,
            y=y,
            layer=normalize_text(row.get("图层")),
            object_type=normalize_text(row.get("对象类型")),
            handle=normalize_text(row.get("句柄")),
        )
        if name in cabinets:
            first = cabinets[name]
            where = f"（CAD句柄 {first.handle} 与 {cabinet.handle}）" if first.handle and cabinet.handle else ""
            warnings.append(f"柜子名称重复：{name}{where}，已采用第一条坐标")
            continue
        cabinets[name] = cabinet
    return cabinets, warnings


def load_segments(data_dir: Path, params: dict[str, float]) -> list[Segment]:
    segments: list[Segment] = []
    unit = params["CAD每米单位"]
    for i, row in enumerate(read_csv_dicts(data_csv_path(data_dir, "路径线段.csv", "route_segments.csv")), start=1):
        ax = parse_float(row.get("起点X"))
        ay = parse_float(row.get("起点Y"))
        bx = parse_float(row.get("终点X"))
        by = parse_float(row.get("终点Y"))
        if ax is None or ay is None or bx is None or by is None:
            continue
        layer = normalize_text(row.get("图层"))
        floor = normalize_floor(row.get("楼层")) or infer_floor_from_layer(layer)
        length_units = parse_float(row.get("长度_CAD单位"))
        chord = math.hypot(bx - ax, by - ay)
        if length_units is None or length_units <= 0:
            length_units = chord
        if length_units <= 0:
            continue
        segments.append(
            Segment(
                idx=i,
                route_id=normalize_text(row.get("路径编号")) or f"R{i}",
                floor=floor,
                ax=ax,
                ay=ay,
                bx=bx,
                by=by,
                length_units=length_units,
                length_m=length_units / unit,
                layer=layer,
                handle=normalize_text(row.get("句柄")),
                section_no=normalize_text(row.get("段号")),
            )
        )
    return segments


def normalize_shaft_base_id(value: str) -> str:
    """竖井编号去掉楼层后缀得到配对名：ZJ1_1F / ZJ1-2F / ZJ1_B1F / 竖井甲二层 → ZJ1 / 竖井甲。"""
    text = normalize_text(value)
    text = re.sub(r"[_\-\s]*(B\d+F?|\d+F|F\d+)$", "", text, flags=re.IGNORECASE)
    text = re.sub(r"(地下|负)?[一二三四五六七八九十]+[层楼]$", "", text)
    return text or "竖井"


def shaft_suffix_floor(shaft_id: str) -> str:
    """从竖井编号尾部解析楼层标注，例如 ZJ1_1F -> 1F；无标注返回空。"""
    match = re.search(r"((?<![A-Za-z])B\d+F?|\d+F|F\d+|(?:地下|负)?[一二三四五六七八九十]+[层楼])$", normalize_text(shaft_id), flags=re.IGNORECASE)
    return normalize_floor(match.group(1)) if match else ""


def load_shafts(data_dir: Path, params: dict[str, float]) -> list[Shaft]:
    shafts: list[Shaft] = []
    default_height = params["竖井高度"]
    for row in read_csv_dicts(data_csv_path(data_dir, "竖井.csv", "shafts.csv")):
        shaft_id = normalize_text(row.get("竖井编号"))
        x = parse_float(row.get("X"))
        y = parse_float(row.get("Y"))
        if not shaft_id or x is None or y is None:
            continue
        layer = normalize_text(row.get("图层"))
        floor = normalize_floor(row.get("楼层")) or infer_floor_from_layer(layer)
        height = parse_float(row.get("高度"), default_height) or default_height
        shafts.append(
            Shaft(
                shaft_id=shaft_id,
                base_id=normalize_shaft_base_id(shaft_id),
                floor=floor,
                x=x,
                y=y,
                height_m=height,
                layer=layer,
                object_type=normalize_text(row.get("对象类型")),
                handle=normalize_text(row.get("句柄")),
            )
        )
    return shafts


def load_rooms(data_dir: Path) -> tuple[list[Room], list[str]]:
    """读取房间边界顶点；同层重名或无效多边形不会参与计算。"""
    rows = read_csv_dicts(data_csv_path(data_dir, "房间范围.csv", "rooms.csv"))
    grouped: dict[tuple[str, str], list[tuple[int, float, float, str, str]]] = defaultdict(list)
    warnings: list[str] = []
    for row_no, row in enumerate(rows, start=2):
        name = normalize_text(row.get("房间名称"))
        layer = normalize_text(row.get("图层"))
        floor = normalize_floor(row.get("楼层")) or infer_floor_from_layer(layer)
        handle = normalize_text(row.get("句柄"))
        vertex_no = parse_float(row.get("顶点序号"))
        x = parse_float(row.get("X"))
        y = parse_float(row.get("Y"))
        if not name or not floor or x is None or y is None:
            warnings.append(f"房间范围第 {row_no} 行无效：需要房间名称、楼层、X、Y")
            continue
        key = (floor, handle or f"名称:{name}")
        grouped[key].append((int(vertex_no or row_no), x, y, name, layer))

    candidates: list[Room] = []
    for (floor, handle_key), points in grouped.items():
        points.sort(key=lambda item: item[0])
        name = points[0][3]
        layer = points[0][4]
        if any(item[3] != name for item in points):
            warnings.append(f"房间范围句柄 {handle_key} 包含多个房间名称，已忽略")
            continue
        vertices: list[tuple[float, float]] = []
        for _, x, y, _, _ in points:
            point = (x, y)
            if not vertices or point != vertices[-1]:
                vertices.append(point)
        if len(vertices) > 1 and vertices[0] == vertices[-1]:
            vertices.pop()
        if len(vertices) < 3 or abs(polygon_signed_area(vertices)) <= 1e-9:
            warnings.append(f"房间范围无效：{floor} {name} 至少需要三个不共线顶点，已忽略")
            continue
        candidates.append(Room(name=name, floor=floor, vertices=vertices, layer=layer, handle="" if handle_key.startswith("名称:") else handle_key))

    name_counts: dict[tuple[str, str], int] = defaultdict(int)
    for room in candidates:
        name_counts[(room.floor, room.name)] += 1
    rooms: list[Room] = []
    for room in candidates:
        if name_counts[(room.floor, room.name)] > 1:
            warnings.append(f"房间名称同层重复：{room.floor} {room.name}，相关范围已忽略")
            continue
        rooms.append(room)
    return rooms, list(dict.fromkeys(warnings))


def load_aliases(data_dir: Path) -> tuple[dict[str, str], list[str]]:
    """读取柜名别名表：清册名称 -> CAD名称。文件不存在时返回空映射。"""
    aliases: dict[str, str] = {}
    warnings: list[str] = []
    for row in read_csv_dicts(data_dir / "柜名别名.csv"):
        source = normalize_cabinet_name(row.get("清册名称"))
        target = normalize_cabinet_name(row.get("CAD名称"))
        if not source or not target:
            continue
        if source == target:
            continue
        if source in aliases:
            if aliases[source] != target:
                warnings.append(f"柜名别名重复且指向不同CAD名称：{source}，已采用第一条")
            continue
        aliases[source] = target
    return aliases, warnings


def apply_aliases(cabinets: dict[str, Cabinet], aliases: dict[str, str]) -> list[str]:
    """把别名注册进柜子表，使清册名称能直接命中CAD柜子坐标。"""
    warnings: list[str] = []
    for source, target in aliases.items():
        cabinet = cabinets.get(target)
        if cabinet is None:
            warnings.append(f"柜名别名目标在 柜子坐标.csv 中不存在：{source} -> {target}")
            continue
        existing = cabinets.get(source)
        if existing is not None and existing is not cabinet:
            warnings.append(f"柜名别名与已有柜子同名，忽略别名：{source} -> {target}")
            continue
        cabinets[source] = cabinet
    return warnings


def apply_case_insensitive_matches(cabinets: dict[str, Cabinet], names: list[str]) -> list[str]:
    """清册柜名与 CAD 柜名只差英文大小写（10KV / 10kV）且唯一对应时自动匹配，并在问题清单里说明。"""
    by_folded: dict[str, list[str]] = defaultdict(list)
    for key in cabinets:
        by_folded[key.casefold()].append(key)
    notes: list[str] = []
    for name in names:
        if not name or name in cabinets:
            continue
        matches = by_folded.get(name.casefold(), [])
        targets = {id(cabinets[key]) for key in matches}
        if len(targets) == 1:
            cabinets[name] = cabinets[matches[0]]
            notes.append(f"提示：柜名只差英文大小写，已自动匹配：{name} -> {matches[0]}")
    return notes


def load_rule_overrides(data_dir: Path) -> tuple[dict[tuple[str, str], dict[str, Any]], list[str]]:
    """读取强制规则表：按起点+终点覆盖自动判断的修正量。文件不存在时返回空映射。"""
    overrides: dict[tuple[str, str], dict[str, Any]] = {}
    warnings: list[str] = []
    for row in read_csv_dicts(data_dir / "强制规则.csv"):
        start = normalize_cabinet_name(row.get("起点"))
        end = normalize_cabinet_name(row.get("终点"))
        extra = parse_float(row.get("修正量"))
        if not start and not end:
            continue
        if not start or not end or extra is None:
            warnings.append(f"强制规则无效（需要起点、终点和数字修正量）：{start or '?'} -> {end or '?'}")
            continue
        if (start, end) in overrides or (end, start) in overrides:
            warnings.append(f"强制规则重复：{start} -> {end}，已采用第一条")
            continue
        overrides[(start, end)] = {"extra": extra, "note": normalize_text(row.get("备注"))}
    return overrides, warnings


def find_header_row(rows: list[tuple[Any, ...]], required: tuple[str, ...] = REQUIRED_HEADERS) -> tuple[int, dict[str, int]] | None:
    """在前几行里找包含全部必需表头的那一行，返回（行号, {表头: 列号}），行列号都从 1 开始。"""
    for row_no, values in enumerate(rows, start=1):
        headers: dict[str, int] = {}
        for col, value in enumerate(values, start=1):
            key = normalize_header(value)
            if key and key not in headers:
                headers[key] = col
        if all(name in headers for name in required):
            return row_no, headers
    return None


def _merged_value_lookup(ws: Any, columns: set[int]) -> dict[tuple[int, int], Any]:
    """纵向合并的单元格（同一起点柜的多条电缆常把起点合并成一格）只有最上面一格有值，把值铺到下面各行。

    横向合并（分组标题行“一、控制电缆”跨好几列）不铺开，否则标题会被当成起点/终点。
    """
    lookup: dict[tuple[int, int], Any] = {}
    for merged in ws.merged_cells.ranges:
        if merged.min_col != merged.max_col or merged.min_col not in columns:
            continue
        value = ws.cell(merged.min_row, merged.min_col).value
        for row in range(merged.min_row + 1, merged.max_row + 1):
            lookup[(row, merged.min_col)] = value
    return lookup


def load_workbook_rows(workbook_path: Path) -> tuple[Any, Any, dict[str, int], list[dict[str, Any]], list[str]]:
    wb = openpyxl.load_workbook(workbook_path)
    ws = wb.active
    head_rows = [
        tuple(cell.value for cell in row)
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, HEADER_SCAN_ROWS))
    ]
    found = find_header_row(head_rows)
    if found is None:
        first = find_header_row(head_rows[:1], ()) or (1, {})
        missing = [h for h in REQUIRED_HEADERS if h not in first[1]]
        raise ValueError(f"清册缺少表头：{', '.join(missing)}（已在前 {HEADER_SCAN_ROWS} 行里查找）")
    header_row, headers = found
    # 手工长度是公式时读 Excel 上次保存的计算结果；列号取插列之前的原始位置
    manual_col = headers["电缆长度"]
    model_col = next((headers[name] for name in MODEL_HEADERS if name in headers), None)
    spec_col = next((headers[name] for name in SPEC_HEADERS if name in headers and headers[name] != model_col), None)
    try:
        values_ws = openpyxl.load_workbook(workbook_path, data_only=True)[ws.title]
    except Exception:
        values_ws = None
    key_cols = {headers["电缆编号"], headers["起点"], headers["终点"]} | {col for col in (model_col, spec_col) if col}
    merged = _merged_value_lookup(ws, key_cols)
    original_cols = {"电缆编号": headers["电缆编号"], "起点": headers["起点"], "终点": headers["终点"]}
    # 多行合并表头的续行、以及横向说明行不能作为电缆数据；只检查关键列，
    # 型号/备注等列的合法合并不会影响起终点行的保留。
    header_continuations: set[int] = set()
    horizontal_key_merges: dict[int, list[tuple[int, int]]] = defaultdict(list)
    for area in ws.merged_cells.ranges:
        title = normalize_header(ws.cell(area.min_row, area.min_col).value)
        if title in REQUIRED_HEADERS and headers[title] == area.min_col and area.min_row >= header_row:
            header_continuations.update(range(area.min_row, area.max_row + 1))
        if area.min_col < area.max_col and any(area.min_col <= col <= area.max_col for col in original_cols.values()):
            for row_no in range(max(header_row + 1, area.min_row), area.max_row + 1):
                horizontal_key_merges[row_no].append((area.min_col, area.max_col))

    def raw(row_no: int, col: int | None) -> Any:
        if col is None:
            return None
        if (row_no, col) in merged:
            return merged[(row_no, col)]
        return ws.cell(row_no, col).value

    # 先按原始列号把数据读出来，再补建列（插列会让右侧单元格整体移位）
    rows: list[dict[str, Any]] = []
    names: set[str] = set()
    for row_no in range(header_row + 1, ws.max_row + 1):
        if row_no in header_continuations:
            continue
        cable_id = normalize_text(raw(row_no, original_cols["电缆编号"]))
        start = normalize_cabinet_name(raw(row_no, original_cols["起点"]))
        end = normalize_cabinet_name(raw(row_no, original_cols["终点"]))
        if not cable_id and not start and not end:
            continue
        # 分页清册常在数据区重复表头，甚至只重复起点/终点两列。
        if normalize_header(cable_id) == "电缆编号" or (start == "起点" and end == "终点"):
            continue
        is_note = False
        for left, right in horizontal_key_merges.get(row_no, ()):
            contains_id = left <= original_cols["电缆编号"] <= right
            covered_endpoints = [
                value for name, value in (("起点", start), ("终点", end))
                if left <= original_cols[name] <= right
            ]
            if (contains_id and covered_endpoints) or (not cable_id and any(covered_endpoints)):
                is_note = True
                break
        if is_note:
            continue
        if start:
            names.add(start)
        if end:
            names.add(end)
        manual = ws.cell(row_no, manual_col).value
        if isinstance(manual, str) and manual.startswith("=") and values_ws is not None:
            manual = values_ws.cell(row_no, manual_col).value
        parsed_manual = parse_length(manual)
        model = " ".join(text for text in (normalize_text(raw(row_no, model_col)), normalize_text(raw(row_no, spec_col))) if text)
        rows.append(
            {
                "row_no": row_no,
                "电缆编号": cable_id,
                "起点": start,
                "终点": end,
                "型号": model,
                "手工长度": parsed_manual if parsed_manual is not None else manual,
            }
        )

    if "自动统计" not in headers:
        auto_col = max(headers.values()) + 1
        ws.cell(header_row, auto_col).value = "自动统计"
        headers["自动统计"] = auto_col
    if "向上取整" not in headers:
        ceil_col = headers["自动统计"] + 1
        if any(col >= ceil_col for col in headers.values()):
            # 「自动统计」右侧还有别的列，插入新列并顺移已记录的列号
            insert_column(ws, ceil_col)
            headers = {name: (col + 1 if col >= ceil_col else col) for name, col in headers.items()}
        ws.cell(header_row, ceil_col).value = "向上取整"
        headers["向上取整"] = ceil_col
    return wb, ws, headers, rows, sorted(names)


_CELL_REF = re.compile(r"(?<![A-Za-z0-9_.$])(\$?)([A-Z]{1,3})(\$?)(\d+)?(?![A-Za-z0-9_(!])")


def _shift_ref(ref: str, from_col: int, sheet_title: str) -> str:
    """把单个引用里 ≥ from_col 的列右移一列；指向其他工作表的引用不动。"""
    sheet, bang, body = ref.rpartition("!")
    if bang and sheet.strip("'") != sheet_title:
        return ref

    def repl(match: re.Match[str]) -> str:
        col = column_index_from_string(match.group(2))
        if col >= from_col:
            col += 1
        return f"{match.group(1)}{get_column_letter(col)}{match.group(3)}{match.group(4) or ''}"

    return f"{sheet}{bang}{_CELL_REF.sub(repl, body)}"


def shift_formula_columns(formula: str, from_col: int, sheet_title: str) -> str:
    """公式里引用本表第 from_col 列及其右侧的地址全部右移一列（插列后保持公式指向不变）。"""
    tokenizer = Tokenizer(formula)
    parts = ["="]
    for token in tokenizer.items:
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            parts.append(_shift_ref(token.value, from_col, sheet_title))
        else:
            parts.append(token.value)
    return "".join(parts)


def insert_column(ws: Any, col: int) -> None:
    """在 col 处插入一列；openpyxl 的 insert_cols 不管合并单元格、列宽和公式引用，这里一并顺移。"""
    merged = [(r.min_row, r.min_col, r.max_row, r.max_col) for r in ws.merged_cells.ranges]
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    widths = {
        column_index_from_string(letter): dim.width
        for letter, dim in ws.column_dimensions.items()
        if dim.width
    }
    ws.insert_cols(col)
    for min_row, min_col, max_row, max_col in merged:
        if min_col >= col:
            min_col, max_col = min_col + 1, max_col + 1
        elif max_col >= col:
            max_col += 1
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
    for index in sorted(widths, reverse=True):
        if index >= col:
            ws.column_dimensions[get_column_letter(index + 1)].width = widths[index]
    ws.column_dimensions[get_column_letter(col)].width = 10
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                try:
                    cell.value = shift_formula_columns(cell.value, col, ws.title)
                except Exception:
                    pass
